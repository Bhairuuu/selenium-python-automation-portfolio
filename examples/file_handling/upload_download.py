import argparse
from pathlib import Path

import openpyxl
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as expected
from selenium.webdriver.support.ui import WebDriverWait

PRACTICE_URL = "https://rahulshettyacademy.com/upload-download-test/index.html"


def update_price(workbook_path: Path, fruit_name: str, new_price: str) -> None:
    workbook = openpyxl.load_workbook(workbook_path)
    sheet = workbook.active
    headers = {cell.value: cell.column for cell in sheet[1]}

    if "price" not in headers:
        raise ValueError("The worksheet does not contain a 'price' column.")

    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row=row, column=1).value == fruit_name:
            sheet.cell(row=row, column=headers["price"]).value = new_price
            workbook.save(workbook_path)
            return

    raise ValueError(f"Fruit '{fruit_name}' was not found in the worksheet.")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Upload an edited product-price workbook to the training site."
    )
    parser.add_argument("workbook", type=Path, help="Downloaded workbook to update.")
    parser.add_argument("--fruit", default="Apple")
    parser.add_argument("--price", default="990")
    arguments = parser.parse_args()

    workbook_path = arguments.workbook.resolve()
    update_price(workbook_path, arguments.fruit, arguments.price)

    driver = webdriver.Chrome()
    wait = WebDriverWait(driver, 10)
    try:
        driver.get(PRACTICE_URL)
        driver.find_element(By.CSS_SELECTOR, "input[type='file']").send_keys(
            str(workbook_path)
        )
        toast = wait.until(
            expected.visibility_of_element_located(
                (By.CSS_SELECTOR, ".Toastify__toast-body div:nth-child(2)")
            )
        )
        assert "uploaded successfully" in toast.text.lower()
    finally:
        driver.quit()


if __name__ == "__main__":
    main()
